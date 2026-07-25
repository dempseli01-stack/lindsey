#!/usr/bin/env python3
"""Lindsey Dempsey — EMEA SDR Leader Challenge PowerPoint.

Built from:
  - Verbatim challenge scripts (speaker notes)
  - Cursor competitor differentiation history
  - Enterprise prep, customer proof (Coinbase / Stripe / PayPal)
  - Official positioning: workflow · context · execution
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

CHARCOAL = RGBColor(0x14, 0x16, 0x18)
SLATE = RGBColor(0x2A, 0x2F, 0x34)
PAPER = RGBColor(0xF4, 0xF5, 0xF2)
INK = RGBColor(0x1A, 0x1C, 0x1E)
MUTED = RGBColor(0x5E, 0x64, 0x6B)
CITRUS = RGBColor(0xD6, 0xF2, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE8, 0xEA, 0xE4)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
OUTPUT = Path(__file__).resolve().parent / "Lindsey_Presentation.pptx"


def set_run(run, *, size, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = anchor
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_bullets(slide, left, top, width, height, lines, *, size=18, color=INK, space_after=11):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = f"•  {line}"
        set_run(run, size=size, color=color, font="Calibri")
    return box


def fill_solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_solid(shape, color)
    return shape


def add_accent(slide, left, top, width=Inches(1.1), height=Inches(0.08)):
    return add_rect(slide, left, top, width, height, CITRUS)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def header_bar(slide, title):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, PAPER)
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.2), CHARCOAL)
    add_textbox(slide, Inches(0.8), Inches(0.32), Inches(11.7), Inches(0.6), title, size=24, bold=True, color=WHITE, font="Georgia")


def title_slide(prs, note):
    s = blank(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, CHARCOAL)
    add_rect(s, 0, Inches(5.7), SLIDE_W, Inches(1.8), SLATE)
    add_accent(s, Inches(0.85), Inches(1.85), width=Inches(1.5))
    add_textbox(s, Inches(0.85), Inches(2.1), Inches(11.5), Inches(1.0), "Up-Level the Machine", size=42, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.85), Inches(3.2), Inches(11.5), Inches(0.5), "Diagnose  →  Protect  →  Experiment  →  Execute", size=22, color=CITRUS)
    add_textbox(s, Inches(0.85), Inches(3.85), Inches(11.5), Inches(0.5), "Player-coach plan for a high-performing EMEA SDR org", size=18, color=SOFT)
    add_textbox(s, Inches(0.85), Inches(6.1), Inches(11.5), Inches(0.4), "Lindsey Dempsey  ·  EMEA SDR Leader Candidate", size=18, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.85), Inches(6.55), Inches(11.5), Inches(0.35), "Cursor wins on workflow · context · execution", size=14, color=SOFT)
    notes(s, note)
    return s


def section_slide(prs, part, title, note):
    s = blank(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, CHARCOAL)
    add_accent(s, Inches(0.85), Inches(2.5), width=Inches(1.4))
    add_textbox(s, Inches(0.85), Inches(2.7), Inches(11), Inches(0.4), part.upper(), size=14, bold=True, color=CITRUS)
    add_textbox(s, Inches(0.85), Inches(3.2), Inches(11.5), Inches(1.2), title, size=32, bold=True, color=WHITE, font="Georgia")
    notes(s, note)
    return s


def content_slide(prs, title, bullets, note, *, size=18):
    s = blank(prs)
    header_bar(s, title)
    add_bullets(s, Inches(0.85), Inches(1.55), Inches(11.5), Inches(5.0), bullets, size=size, space_after=12)
    notes(s, note)
    return s


def two_col(prs, title, left_title, left_items, right_title, right_items, note, *, size=15):
    s = blank(prs)
    header_bar(s, title)
    add_rect(s, Inches(0.65), Inches(1.45), Inches(5.85), Inches(5.2), WHITE)
    add_accent(s, Inches(0.95), Inches(1.7), width=Inches(0.85))
    add_textbox(s, Inches(0.95), Inches(1.95), Inches(5.2), Inches(0.4), left_title, size=17, bold=True, color=CHARCOAL, font="Georgia")
    add_bullets(s, Inches(0.95), Inches(2.5), Inches(5.2), Inches(3.8), left_items, size=size, space_after=8)

    add_rect(s, Inches(6.8), Inches(1.45), Inches(5.85), Inches(5.2), WHITE)
    add_accent(s, Inches(7.1), Inches(1.7), width=Inches(0.85))
    add_textbox(s, Inches(7.1), Inches(1.95), Inches(5.2), Inches(0.4), right_title, size=17, bold=True, color=CHARCOAL, font="Georgia")
    add_bullets(s, Inches(7.1), Inches(2.5), Inches(5.2), Inches(3.8), right_items, size=size, space_after=8)
    notes(s, note)
    return s


def market_slide(prs, note):
    s = blank(prs)
    header_bar(s, "Market context — agentic coding is a two-layer race")

    add_rect(s, Inches(0.65), Inches(1.45), Inches(5.85), Inches(3.35), WHITE)
    add_rect(s, Inches(0.65), Inches(1.45), Inches(5.85), Inches(0.5), CHARCOAL)
    add_textbox(s, Inches(0.9), Inches(1.52), Inches(5.3), Inches(0.4), "Layer 1 — Harness / product", size=15, bold=True, color=CITRUS, font="Georgia")
    add_bullets(s, Inches(0.9), Inches(2.15), Inches(5.3), Inches(2.4), [
        "Plans, tools, repo navigation, developer workflow",
        "What buyers feel in the first conversation",
        "Cursor’s established commercial strength",
        "Sell the system they can feel today",
    ], size=14, space_after=7)

    add_rect(s, Inches(6.8), Inches(1.45), Inches(5.85), Inches(3.35), WHITE)
    add_rect(s, Inches(6.8), Inches(1.45), Inches(5.85), Inches(0.5), SLATE)
    add_textbox(s, Inches(7.05), Inches(1.52), Inches(5.3), Inches(0.4), "Layer 2 — Model", size=15, bold=True, color=WHITE, font="Georgia")
    add_bullets(s, Inches(7.05), Inches(2.15), Inches(5.3), Inches(2.4), [
        "Reasoning + long-horizon coding quality",
        "Composer = Cursor’s coding-specialist path",
        "Best model changes over time",
        "Frontier labs = suppliers and competitors",
    ], size=14, space_after=7)

    add_rect(s, Inches(0.65), Inches(5.0), Inches(11.95), Inches(1.55), CHARCOAL)
    add_textbox(s, Inches(0.9), Inches(5.15), Inches(11.4), Inches(0.35), "Competitive shapes (from differentiation history)", size=13, bold=True, color=CITRUS)
    add_textbox(
        s,
        Inches(0.9),
        Inches(5.55),
        Inches(11.4),
        Inches(0.8),
        "Cursor = AI-native IDE + enterprise platform   ·   Copilot = plugin / GitHub ecosystem   ·   Claude Code = terminal-first agent   ·   Windsurf = AI-IDE challenger (Cascade)   ·   Often the real competitor = ungoverned shadow AI",
        size=13,
        color=SOFT,
    )
    notes(s, note)
    return s


def positioning_slide(prs, note):
    s = blank(prs)
    header_bar(s, "Cursor positioning — workflow · context · execution")

    add_rect(s, Inches(0.65), Inches(1.4), Inches(11.95), Inches(1.0), CHARCOAL)
    add_textbox(
        s,
        Inches(0.9),
        Inches(1.55),
        Inches(11.4),
        Inches(0.7),
        "Leading agentic coding platform — shared AI layer on the codebase that connects developers, tools, and agents so teams get more from frontier models.",
        size=15,
        color=SOFT,
    )

    surfaces = [("IDE", "Desktop agentic coding"), ("CLI", "Terminal / any workflow"), ("Automations", "Event-driven agents"), ("Cloud Agents", "Async / browser / mobile")]
    x = Inches(0.65)
    for label, sub in surfaces:
        add_rect(s, x, Inches(2.65), Inches(2.85), Inches(1.25), WHITE)
        add_accent(s, x + Inches(0.2), Inches(2.85), width=Inches(0.65))
        add_textbox(s, x + Inches(0.2), Inches(3.05), Inches(2.45), Inches(0.3), label, size=15, bold=True, color=CHARCOAL, font="Georgia")
        add_textbox(s, x + Inches(0.2), Inches(3.4), Inches(2.45), Inches(0.3), sub, size=12, color=MUTED)
        x += Inches(3.05)

    pillars = [
        ("Workflow", "How work actually gets done across the SDLC"),
        ("Context", "Codebase, dependencies, team systems"),
        ("Execution", "Plan → build → review → ship with less friction"),
    ]
    x = Inches(0.65)
    for title, body in pillars:
        add_rect(s, x, Inches(4.2), Inches(3.9), Inches(1.35), WHITE)
        add_textbox(s, x + Inches(0.25), Inches(4.4), Inches(3.4), Inches(0.35), title, size=16, bold=True, color=CHARCOAL, font="Georgia")
        add_textbox(s, x + Inches(0.25), Inches(4.85), Inches(3.4), Inches(0.5), body, size=13, color=MUTED)
        x += Inches(4.1)

    add_textbox(
        s,
        Inches(0.85),
        Inches(5.8),
        Inches(11.5),
        Inches(0.85),
        "Enterprise 3 buckets:  Productivity (agents + multi-file)  ·  Governance (SSO / SCIM / audit / admin)  ·  Privacy & security (Privacy Mode, SOC 2 Type II)\nProof points: Coinbase · Stripe · PayPal — speed + governed adoption, not autocomplete vanity.",
        size=13,
        color=MUTED,
    )
    notes(s, note)
    return s


def differentiators_slide(prs, note):
    s = blank(prs)
    header_bar(s, "Four differentiators — diagnose first, lead with one wedge")
    cards = [
        ("01", "Model neutrality", "Best model changes. Flexibility over single-vendor lock-in. Labs can be suppliers and competitors."),
        ("02", "Large codebase / harness", "Not just wrapping a model — better finding + using context in messy enterprise repos."),
        ("03", "Faster time to value", "Useful quickly without every developer inventing a custom setup. Adoption + consistency."),
        ("04", "Platform across the SDLC", "Plan, write, review, debug, iterate — Bugbot, Agent Review, Automations. Better execution."),
    ]
    positions = [(Inches(0.65), Inches(1.45)), (Inches(6.8), Inches(1.45)), (Inches(0.65), Inches(4.05)), (Inches(6.8), Inches(4.05))]
    for (left, top), (num, title, body) in zip(positions, cards):
        add_rect(s, left, top, Inches(5.85), Inches(2.35), WHITE)
        add_accent(s, left + Inches(0.25), top + Inches(0.25), width=Inches(0.75))
        add_textbox(s, left + Inches(0.25), top + Inches(0.45), Inches(0.9), Inches(0.35), num, size=16, bold=True, color=CHARCOAL, font="Georgia")
        add_textbox(s, left + Inches(1.15), top + Inches(0.45), Inches(4.3), Inches(0.35), title, size=16, bold=True, color=CHARCOAL, font="Georgia")
        add_textbox(s, left + Inches(0.25), top + Inches(1.05), Inches(5.3), Inches(1.0), body, size=14, color=MUTED)
    notes(s, note)
    return s


def competitive_battlecard(prs, note):
    s = blank(prs)
    header_bar(s, "Why we win — competition differentiation")

    add_rect(s, Inches(0.55), Inches(1.35), Inches(12.15), Inches(0.8), CHARCOAL)
    add_textbox(
        s,
        Inches(0.75),
        Inches(1.45),
        Inches(11.7),
        Inches(0.6),
        "We don’t win by claiming exclusive smarter intelligence forever. We win by applying intelligence better — workflow, codebase context, integrated SDLC — with model neutrality.",
        size=13,
        color=SOFT,
    )

    cols = [
        (Inches(0.55), "Cursor", CITRUS, CHARCOAL, [
            "AI-native IDE + enterprise platform",
            "Snippet help → task completion",
            "Shared AI layer on the codebase",
            "IDE · CLI · Automations · Cloud Agents",
            "Multi-model + Composer path",
            "Bugbot / review clears PR bottleneck",
        ]),
        (Inches(3.7), "vs GitHub Copilot", WHITE, SLATE, [
            "Helps developers write faster",
            "Often felt as local / snippet help",
            "Wins: GitHub ecosystem + friction",
            "Our move: task completion",
            "Our move: broader codebase context",
            "Our move: integrated experience",
        ]),
        (Inches(6.85), "vs Claude Code", WHITE, SLATE, [
            "Strong — don’t trash-talk it",
            "Terminal-first agent shape",
            "Not a “smarter model” war",
            "Our edge: model neutrality",
            "Our edge: SDLC integration",
            "Our edge: team time-to-value",
        ]),
        (Inches(10.0), "Also in market", WHITE, SLATE, [
            "Windsurf: Cascade agent UX",
            "Codex: OpenAI async agents",
            "Shadow AI: real enterprise rival",
            "Win = governed standardization",
            "Proof: Coinbase / Stripe / PayPal",
            "Outcomes > autocomplete vanity",
        ]),
    ]
    for left, title, title_color, bar, items in cols:
        add_rect(s, left, Inches(2.35), Inches(3.0), Inches(4.25), WHITE)
        add_rect(s, left, Inches(2.35), Inches(3.0), Inches(0.5), bar)
        add_textbox(s, left + Inches(0.12), Inches(2.42), Inches(2.75), Inches(0.38), title, size=12, bold=True, color=title_color if bar != WHITE else CHARCOAL, font="Georgia")
        add_bullets(s, left + Inches(0.12), Inches(3.05), Inches(2.75), Inches(3.3), items, size=11, space_after=5)
    notes(s, note)
    return s


def numbered_slide(prs, title, items, note):
    s = blank(prs)
    header_bar(s, title)
    top = Inches(1.45)
    for num, heading, body in items:
        add_textbox(s, Inches(0.85), top, Inches(0.7), Inches(0.35), num, size=17, bold=True, color=CHARCOAL, font="Georgia")
        add_textbox(s, Inches(1.55), top, Inches(10.8), Inches(0.32), heading, size=16, bold=True, color=CHARCOAL)
        add_textbox(s, Inches(1.55), top + Inches(0.32), Inches(10.8), Inches(0.5), body, size=14, color=MUTED)
        top += Inches(1.15)
    notes(s, note)
    return s


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1
    title_slide(
        prs,
        """Hi everyone — I’m Lindsey Dempsey. Thank you for having me.

Today I’m going to walk through how I would up-level the EMEA SDR machine — not rebuild it. My framing is simple: diagnose, protect, experiment, and execute, as a hands-on player-coach.

I’ll keep this practical and leave room for discussion and a live pitch.""",
    )

    # 2
    content_slide(
        prs,
        "Thesis: don’t rebuild a working machine — raise its ceiling",
        [
            "EMEA already high-performing: AI-augmented workflows, clear ICP, eight-figure monthly pipeline",
            "Job = find real constraints, protect what works, run a few high-leverage experiments",
            "Every SDR fluent in how Cursor wins — workflow, context, execution",
            "Lead as a player-coach: trains on time, inspection over activity theater",
            "Success = better conversation quality + higher AE acceptance — not a reorg",
        ],
        """My thesis is this: EMEA is already high-performing. You have AI-augmented workflows, a clear ICP, and you’re generating eight-figure qualified pipeline every month. So the question isn’t “how do I stand something up from zero?” It’s “given a function that already works, how do I make it materially better?”

My job would be to find the real constraints, protect what’s working, and run a small number of high-leverage experiments.

Every SDR on the team should be fluent in how Cursor wins — on workflow, context, and execution.

And I would lead as a player-coach. Trains on time. Inspection over activity theater. Success for me is better conversation quality and higher AE acceptance — not reinventing the org chart.""",
    )

    # 3
    content_slide(
        prs,
        "Agenda",
        [
            "Part 1 (20 min) — First 15 / 30 / 60: diagnose, protect, gains, competitive context, two experiments",
            "Part 2 (12 min) — Fundamentals: hire, coach, run the team as a player-coach",
            "Part 3 (10 min) — Craft: live pitch, persona outreach, why Cursor wins (Copilot + Claude Code)",
            "Clarity over polish — happy to go live earlier on the pitch",
        ],
        """Here’s how I’ll use the time.

Part one is the first fifteen, thirty, and sixty days — diagnosis, what I’d protect, where the gains are, and two experiments.

Part two is fundamentals — how I hire, coach, and run the team day to day as a player-coach.

Part three is craft — how I’d pitch, how we cut through noise by persona, and why Cursor wins, including Copilot and Claude Code.

Clarity over polish. If you want me to go live earlier on the pitch, I’m happy to.""",
    )

    # 4
    section_slide(prs, "Part 1", "How I’d up-level EMEA in the first sixty days", "Part one — how I’d up-level EMEA in the first sixty days.")

    # 5
    two_col(
        prs,
        "Days 1–15: diagnosis only",
        "Data",
        [
            "Full funnel by country, persona, channel, rep",
            "AE acceptance + stage progression (meeting quality truth serum)",
            "Where meetings die: brush-offs, Copilot, Claude Code, security, no next step",
            "Account quality: 250+ engineers, AI tooling already in place",
            "Call sample: Cursor wedge vs feature dumping",
        ],
        "People + constraints",
        [
            "Interview every SDR (top + bottom), AEs, Growth, SE/Solutions, sales leadership",
            "Define “up-level” in numbers with leadership",
            "Constraint hunt: connect → conversation",
            "Meeting → AE acceptance",
            "Competitive framing",
        ],
        """Days one through fifteen are diagnosis only. I will not make big changes before I understand the machine.

On data, I’d pull the full funnel by country, persona, channel, and rep. I’d look hard at AE acceptance and stage progression — because that’s the truth serum on meeting quality. I’d also look at where meetings die: brush-offs, Copilot objections, Claude Code objections, security pushback, or no clear next step. I’d check account quality signals like two-hundred-fifty-plus engineers and AI tooling already in place. And I’d sample calls to see what percent of conversations are using the right Cursor wedge versus feature dumping.

On people, I’d interview every SDR — especially top and bottom performers — AEs on what makes a meeting great versus junk, Growth on automation and feedback loops, Solutions or SE on where reps get technically fuzzy, and sales leadership on what “up-level” actually means in numbers.

My constraint hunt would focus on three places strong teams usually leak: connect-to-conversation, meeting-to-AE-acceptance, and competitive framing.""",
    )

    # 6
    content_slide(
        prs,
        "Protect what’s already working",
        [
            "Tooling + AI research automation stack",
            "Current ICP and target-market strategy producing pipeline",
            "Execution cadence that keeps trains on time",
            "Top-performer craft — don’t standardize away what your best people do",
            "AE trust and qualification norms",
            "Explicit: no reorg, rename stages, or full playbook rewrite in month one",
        ],
        """Just as important as what I’d change is what I would protect.

I would protect the existing tooling and AI research automation stack. I would protect the current ICP and target-market strategy that’s already producing pipeline. I would protect the execution cadence that keeps trains on time. I would protect top-performer craft — I won’t standardize away what your best people are doing. And I would protect AE trust and qualification norms.

I want to be explicit: I will not reorg, rename stages, or rewrite the full playbook in month one. That’s how leaders break high-performing teams.""",
    )

    # 7
    content_slide(
        prs,
        "Where the next gains come from (quality compounding)",
        [
            "Conversation quality — relevance speed + correct wedge for the buyer signal",
            "Qualification bar — fewer junk meetings, higher AE acceptance",
            "Competitive fluency — clean Copilot / Claude Code / Windsurf frames (no trash talk)",
            "Earlier multi-threading — Platform + Security on enterprise / regulated accounts",
            "Validate with diagnosis data before scaling changes",
        ],
        """When a team is already strong on volume, the next level usually comes from quality compounding.

First, conversation quality — relevance speed with engineering leaders, and picking the correct wedge for the buyer signal.

Second, the qualification bar — fewer junk meetings, higher AE acceptance.

Third, competitive fluency — clean Copilot and Claude Code frames without trash talk or feature dumps.

Fourth, earlier multi-threading — especially Platform and Security on enterprise and regulated accounts — so opportunities are stronger when they enter the AE motion.

That’s where I’d look first, then validate with the diagnosis data.""",
    )

    # 8
    market_slide(
        prs,
        """Quick market context, because it shapes how we coach messaging.

Agentic coding is now a two-layer race.

Layer one is the harness and product layer — how the system plans, uses tools, navigates repositories, and fits developer workflow. That’s where Cursor has already built commercial strength, and it’s what buyers feel in a first conversation.

Layer two is the model layer — reasoning quality and long-horizon coding performance. Composer is Cursor’s coding-specialist model path, and that layer continues to strengthen.

Model neutrality still matters, because the best model changes over time, and the frontier labs can be both suppliers and competitors.

For EMEA SDRs, the implication is simple: sell the system buyers can feel today — workflow, context, and execution — and stay fluent on where the model layer is going. I won’t turn customer calls into infrastructure or acquisition conversations.""",
    )

    # 9
    positioning_slide(
        prs,
        """This is the spine I’d train every rep on.

Cursor is the leading agentic coding platform that a large share of the Fortune one thousand use to reimagine how engineering teams build software with AI.

It gives teams a shared AI layer that understands their codebase, dependencies, and workflows — connecting developers, tools, and AI agents into one system.

It’s helpful to think of Cursor as a coding agent that can be accessed, coordinated, and controlled through product surfaces like the IDE, CLI, Automations, and Cloud Agents.

It bridges planning, execution, and validation so work moves through the SDLC with less friction.

The summary line I want reps to be able to say is: Cursor helps teams get more out of frontier models by pairing them with better workflow, better context, and a more integrated way to build software.""",
    )

    # 10
    differentiators_slide(
        prs,
        """Cursor’s core differentiation is the combination of four things, and I’d teach SDRs to lead with only one based on the buyer’s pain.

One — model neutrality. Customers want flexibility. The best model changes. They don’t want to be locked to one vendor.

Two — large codebase performance. Enterprise work spans messy, real-world repos. Cursor isn’t just wrapping a model — it improves how context is found and used. That matters more as complexity increases.

Three — faster time to value. Cursor is useful quickly without every developer inventing a custom setup. Enterprises care about adoption and consistency across teams.

Four — platform across the SDLC. Not just code generation — planning, writing, reviewing, debugging, iterating — through things like Bugbot, Agent Review, and Automations. The value is better execution, not just better answers.

If a rep tries to say all four on every call, they’ll lose the room. Diagnose first, then pick the wedge.""",
    )

    # 11
    two_col(
        prs,
        "Days 15–45: two experiments only (keep or kill)",
        "A — Positioning fluency",
        [
            "Hypothesis: connect→conversation + AE credibility rise when every opener maps to one differentiator",
            "Test: 2 SDRs × 2 weeks; dial volume held constant",
            "Measure: connect→conversation, conversation→meeting, AE credibility feedback",
        ],
        "B — Competitive tracks",
        [
            "Track A (Copilot-installed): snippet help → task completion + codebase workflow",
            "Track B (Claude Code pockets): neutrality + SDLC integration, not a model war",
            "Measure: positive replies, meetings, objection handling, opportunity acceptance",
        ],
        """Between day fifteen and forty-five, I’d run two experiments only. Small bets, hard metrics, keep or kill.

Experiment A is positioning fluency. Hypothesis: connect-to-conversation and AE credibility rise when every opener maps to one differentiator by buyer signal. I’d test with two SDRs for two weeks and hold dial volume constant, so we don’t confuse activity with lift. I’d measure connect-to-conversation, conversation-to-meeting, and AE feedback on whether the rep sounded credible.

Experiment B is competitive tracks. Track A for Copilot-installed accounts — move the story from snippet help to task completion and codebase workflow. Track B for Claude Code pockets — neutrality plus SDLC integration, not a model war. Same idea: measure positive replies, meetings, objection handling, and opportunity acceptance.

If it doesn’t move the metric, we kill it. No sacred cows.""",
    )

    # 12
    content_slide(
        prs,
        "Day 60 outcomes + operating loops",
        [
            "Baseline dashboard + clear “constraint of the month”",
            "Experiments concluded with keep-or-kill decisions",
            "Updated hiring bar for next SDR hires",
            "Locked weekly coaching cadence + my own call floor (player-coach)",
            "Weekly triangle with Growth + AE lead: what converted, what failed, one playbook update",
            "Every change gets a metric owner and a review date — trains on time",
        ],
        """By day sixty, I want five things true: a baseline dashboard with a clear “constraint of the month,” experiments concluded with keep-or-kill decisions, an updated hiring bar for the next SDR hires, a locked weekly coaching cadence, and my own player-coach call floor — I’m in the work, not above it.

Cross-functionally, I’d run a weekly triangle with Growth and an AE lead. Agenda is simple: what converted, what failed, and one playbook update only. No thrash.

SDRs execute. Automation helps with research and selection. I inspect quality. Every change gets a metric owner and a review date. That’s how you keep trains on time while still improving.""",
    )

    # 13
    section_slide(prs, "Part 2", "Fundamentals — how I lead day to day as a player-coach", "Part two — fundamentals. How I lead day to day.")

    # 14
    content_slide(
        prs,
        "The ideal AI-era SDR",
        [
            "Judgment — which account and persona deserve the dial",
            "Relevance speed — earn 30 seconds with a technical buyer",
            "Positioning fluency — workflow, context, execution, correct wedge",
            "Competitive clarity — Copilot + Claude Code + Windsurf without trash talk",
            "Security instincts — calm on Privacy Mode and enterprise controls",
            "Execution discipline — trains on time",
            "Bar: diagnose the buyer’s constraint and match Cursor to it — don’t recite features",
        ],
        """In the AI era, a lot of rote research and first-draft personalization can be automated. So the exceptional SDR looks different.

They win on judgment — which account and persona deserve the dial.
Relevance speed — earning thirty seconds with a technical buyer.
Positioning fluency — workflow, context, execution, and the correct wedge.
Competitive clarity — Copilot and Claude Code without trash talk.
Security instincts — calm and grounded on Privacy Mode and enterprise controls.
And execution discipline — trains on time.

My bar is simple: they don’t recite features. They diagnose the buyer’s constraint and match Cursor to it.""",
    )

    # 15
    two_col(
        prs,
        "Hiring + coaching",
        "Hiring bar",
        [
            "Coachable, clear communicators, curious about technical buyers, process discipline, learn fast",
            "Hard nos: feature dumpers, activity theater, can’t take critique, overclaim the product",
            "Assess: mock call (VP Eng / Platform), Copilot + Claude Code objections, 60-sec Cursor pitch, soft security question",
        ],
        "Coaching system",
        [
            "Daily scoreboard + same-day course-correction",
            "Weekly 1:1s — one funnel leak only",
            "Weekly team: one win wire + one talk-track upgrade",
            "Call reviews: opener, wedge, competitive frame, clean ask",
            "Underperformance = stage diagnosis + skill plan + timeline",
        ],
        """On hiring, I’m looking for coachable people who communicate clearly, are curious about technical buyers, have process discipline, and learn fast.

Hard nos: feature dumpers, activity theater without conversion ownership, people who can’t take critique, and anyone who overclaims the product.

I’d assess with a mock call to a VP of Engineering or Platform lead, Copilot and Claude Code objections, a sixty-second Cursor pitch using the official framing, and a soft security question like, “Our CISO won’t allow AI on our code — what do you say?”

On coaching: daily scoreboard visibility and same-day course-correction if someone misses the standard. Weekly one-on-ones focused on one funnel leak only. Weekly team session with one win wire and one talk-track upgrade. Call reviews scored on opener, wedge, competitive frame, and clean ask.

Underperformance gets a stage diagnosis, a skill plan, and a timeline — not vague pressure.""",
    )

    # 16
    section_slide(prs, "Part 3", "Craft and curiosity — happy to go live here", "Part three — craft and curiosity. This is where I’m happiest going live.")

    # 17
    numbered_slide(
        prs,
        "Live pitch spine: open → position → wedge → ask",
        [
            ("01", "Open", "From AI that helps write snippets → AI that helps complete work across the codebase."),
            ("02", "Position", "Agentic coding platform across IDE, CLI, Automations, Cloud Agents — shared AI layer on your codebase."),
            ("03", "Wedge", "Pick one differentiator based on what you heard."),
            ("04", "Ask", "Worth 20 minutes to map leverage vs where workflow still breaks — write, review, or ship?"),
        ],
        """If I’m cold-calling into this room, the spine is open, position, wedge, ask.

Open: the reason I’m calling is engineering teams are moving from AI that helps write snippets to AI that helps complete work across the codebase — and most tools only cover part of that workflow.

Position: Cursor is an agentic coding platform — a coding agent across the IDE, CLI, Automations, and Cloud Agents — with a shared AI layer that understands your codebase and workflows so planning, building, review, and shipping happen with less friction.

Wedge: I pick one differentiator based on what I heard.

Ask: worth twenty minutes to map where you’re getting leverage versus where workflow still breaks — write, review, or ship?

If I get “just send an email,” I don’t collapse. I say: happy to — so I send the useful one: are you more focused on standardizing AI coding, or on the review bottleneck after Copilot?

I’m happy to run this live on one of you right now if you’d like.""",
    )

    # 18
    content_slide(
        prs,
        "Cutting through noise — persona-true outreach",
        [
            "Hands-on developer: real-repo context + daily workflow (avoid corporate ROI theater)",
            "VP of Engineering: shared AI layer, SDLC speed, standardizing teams (avoid feature laundry lists)",
            "Security / Platform: Privacy Mode, admin controls, governed rollout (avoid hype)",
            "What works: trigger-based relevance (Copilot sprawl, AI policy, migrations, DevEx) + one sharp question",
            "What fails: generic “quick chat” spam and fake personalization",
            "Proof to arm reps: Coinbase (agent-first), Stripe (rollout), PayPal (enterprise velocity)",
        ],
        """Technical buyers are saturated on LinkedIn and email, so persona-true outreach matters more than volume tricks.

For a hands-on developer: lead with real-repo context and daily workflow. Avoid corporate ROI theater.

For a VP of Engineering: lead with a shared AI layer, speed across the SDLC, and standardizing teams. Avoid feature laundry lists.

For Security or Platform: lead with Privacy Mode, admin controls, and governed rollout so AI coding doesn’t become shadow IT. Avoid hype.

What still works is trigger-based relevance — Copilot sprawl, AI policy moments, migrations, DevEx ownership — and one sharp question. What fails is generic “quick chat” spam and fake personalization.""",
        size=16,
    )

    # 19
    competitive_battlecard(
        prs,
        """For a skeptical engineering leader, my one-to-two sentence version is this:

Cursor doesn’t win because we claim exclusive access to better intelligence forever. We win because we apply intelligence more effectively — through workflow, codebase context, and an integrated system across the SDLC — with model neutrality as capabilities and economics change.

Versus GitHub Copilot: Copilot helps developers write faster. Cursor is the shift from snippet help to task completion, from local suggestions to broader codebase context, from a point solution to a more integrated development experience.

Versus Claude Code: Claude Code is strong, and I won’t argue it’s weak. Our edge usually isn’t exclusive smarter model access. It’s model neutrality instead of single-provider dependence, stronger integration across the SDLC, better support for large complex codebases, and faster time to value across a broader team.""",
    )

    # 20
    two_col(
        prs,
        "Security instincts (not a certification dump)",
        "Privacy Mode",
        [
            "Customer code is not stored or retained",
            "Not used to train models",
            "Requests are isolated and ephemeral",
            "Enterprise: enforce org-wide so shadow AI doesn’t win",
        ],
        "Enterprise posture",
        [
            "SOC 2 Type II; secure handling of code + metadata",
            "Admin controls, visibility, model / MCP governance",
            "SSO / SCIM / audit logs for regulated buyers",
            "First-call frame: adoption confidence — bring the right partner when deep",
        ],
        """Security and privacy often decide enterprise evaluations, so SDRs need calm instincts, not a certification dump.

On Privacy Mode: customer code is not stored or retained, not used to train models, and requests are isolated and ephemeral.

On enterprise posture: SOC 2, secure handling of code and metadata, admin controls and visibility, and the fact that security requirements influence product design.

In a live cycle, security shows up as adoption confidence — “we can roll this out without creating shadow AI” — not as a slide war on the first call. If it gets deep, I bring in the right technical partner and stay precise.""",
    )

    # 21
    content_slide(
        prs,
        "Close — up-level, don’t rebuild",
        [
            "Diagnose real constraints with data + AE/SDR truth",
            "Protect automation, ICP, and cadence already creating pipeline",
            "Train workflow / context / execution — four wedges, clean competitive frames",
            "Run one or two measured experiments — keep or kill on metrics",
            "Lead as a player-coach so the floor rises with the ceiling",
            "Questions — and happy to pitch one of you live right now",
        ],
        """I’ll close where I started.

In the first sixty days I won’t invent a new EMEA. I’ll diagnose the real constraints with data and AE/SDR truth. I’ll protect the automation, ICP, and cadence already creating pipeline. I’ll train the team on workflow, context, and execution — four wedges, clean competitive frames. I’ll run one or two measured experiments and keep or kill them on metrics. And I’ll lead as a player-coach so the floor rises with the ceiling.

That’s how you up-level a function that’s already working — and keep the trains on time.

I’m Lindsey Dempsey. I’d love your questions — and if useful, I’m ready to pitch one of you live right now.""",
    )

    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
